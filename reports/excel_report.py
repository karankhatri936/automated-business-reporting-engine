"""Excel report generation using OpenPyXL.

Builds a multi-sheet workbook:

1. Dashboard          - title, KPI summary, inventory status, top products, charts
2. Product Data       - processed product dataset as an Excel table
3. Category Analysis  - category-level metrics as an Excel table
4. Product Analysis   - product-level rankings
5. Report Metadata    - API source, generation info and record count

The generator consumes the SAME KPI results produced by
``KPIEngine.calculate_all()`` so Excel and PDF reports stay consistent.
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..config import REPORT_VERSION, OutputConfig, config
from ..utils.exceptions import ExcelReportError
from ..utils.logger import setup_logger

logger = setup_logger(__name__)

# ── Number formats ─────────────────────────────────────────────
CURRENCY = '"$"#,##0.00'
INTEGER = '#,##0'
DECIMAL2 = '0.00'
PERCENT = '0.0"%"'

# ── Style constants ────────────────────────────────────────────
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F3864")
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F3864")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
HEADER_FILL = PatternFill(fill_type="solid", start_color="1F3864", end_color="1F3864")
SECTION_FILL = PatternFill(fill_type="solid", start_color="D6E4F0", end_color="D6E4F0")
ALT_ROW_FILL = PatternFill(fill_type="solid", start_color="F2F7FC", end_color="F2F7FC")

_THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# ── Column specifications per sheet ────────────────────────────
# Each entry: (source column, display header, number format)

PRODUCT_COLUMNS = [
    ("id", "ID", INTEGER),
    ("title", "Title", None),
    ("brand", "Brand", None),
    ("category", "Category", None),
    ("price", "Price", CURRENCY),
    ("rating", "Rating", DECIMAL2),
    ("stock", "Stock", INTEGER),
    ("discountPercentage", "Discount %", PERCENT),
    ("availabilityStatus", "Availability", None),
    ("inventory_value", "Inventory Value", CURRENCY),
    ("stock_status", "Stock Status", None),
]

CATEGORY_COLUMNS = [
    ("category", "Category", None),
    ("product_count", "Products", INTEGER),
    ("average_price", "Average Price", CURRENCY),
    ("average_rating", "Average Rating", DECIMAL2),
    ("total_stock", "Total Stock", INTEGER),
    ("inventory_value", "Inventory Value", CURRENCY),
    ("average_discount", "Average Discount", PERCENT),
]

PRODUCT_ANALYSIS_SECTIONS = [
    (
        "top_by_inventory_value",
        "Top 10 Products by Inventory Value",
        [
            ("id", "ID", INTEGER),
            ("title", "Title", None),
            ("category", "Category", None),
            ("price", "Price", CURRENCY),
            ("stock", "Stock", INTEGER),
            ("inventory_value", "Inventory Value", CURRENCY),
        ],
    ),
    (
        "top_by_price",
        "Top 10 Products by Price",
        [
            ("id", "ID", INTEGER),
            ("title", "Title", None),
            ("category", "Category", None),
            ("price", "Price", CURRENCY),
            ("inventory_value", "Inventory Value", CURRENCY),
        ],
    ),
    (
        "top_by_rating",
        "Top 10 Products by Rating",
        [
            ("id", "ID", INTEGER),
            ("title", "Title", None),
            ("category", "Category", None),
            ("rating", "Rating", DECIMAL2),
            ("inventory_value", "Inventory Value", CURRENCY),
        ],
    ),
    (
        "lowest_stock",
        "10 Products with Lowest Stock",
        [
            ("id", "ID", INTEGER),
            ("title", "Title", None),
            ("category", "Category", None),
            ("stock", "Stock", INTEGER),
            ("inventory_value", "Inventory Value", CURRENCY),
        ],
    ),
    (
        "top_discounted",
        "Top 10 Products by Discount",
        [
            ("id", "ID", INTEGER),
            ("title", "Title", None),
            ("category", "Category", None),
            ("discountPercentage", "Discount %", PERCENT),
            ("inventory_value", "Inventory Value", CURRENCY),
        ],
    ),
]

TABLE_NAMES = {
    "top_by_inventory_value": "TopByInventoryValue",
    "top_by_price": "TopByPrice",
    "top_by_rating": "TopByRating",
    "lowest_stock": "LowestStock",
    "top_discounted": "TopDiscounted",
}

KPI_META = [
    ("total_products", "Total Products", INTEGER),
    ("average_price", "Average Price", CURRENCY),
    ("median_price", "Median Price", CURRENCY),
    ("total_inventory_units", "Total Inventory Units", INTEGER),
    ("average_stock_per_product", "Average Stock per Product", DECIMAL2),
    ("total_inventory_value", "Total Inventory Value", CURRENCY),
    ("average_rating", "Average Rating", DECIMAL2),
    ("average_discount", "Average Discount", PERCENT),
    ("low_stock_count", "Low Stock Products", INTEGER),
]

STOCK_STATUS_ORDER = ["out_of_stock", "low", "medium", "high"]

# ── Small helpers ──────────────────────────────────────────────


def _is_missing(value: Any) -> bool:
    """Return True for None, NaN, NaT or blank strings."""
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _excel_value(value: Any) -> Any:
    """Convert values into types OpenPyXL can write safely."""
    if _is_missing(value):
        return None
    if isinstance(value, (str, int, float, bool, datetime)):
        return value
    if hasattr(value, "item"):  # numpy scalars
        return value.item()
    return value


def _auto_widths(ws, df: pd.DataFrame, col_specs) -> None:
    """Set readable column widths from header/data length (capped)."""
    for i, (col, header, _fmt) in enumerate(col_specs, start=1):
        if col not in df.columns:
            continue
        length = len(str(header))
        for value in df[col].head(200):
            if _is_missing(value):
                continue
            length = max(length, min(len(str(value)), 40))
        ws.column_dimensions[get_column_letter(i)].width = min(length + 3, 50)


def _style_header_row(ws, row: int, ncols: int) -> None:
    """Apply professional header styling to a row of cells."""
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_section_title(ws, row: int, text: str, max_col: int = 2) -> int:
    """Write a shaded section header across ``max_col`` columns; returns next row."""
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=max_col
    )
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    return row + 1


def _write_data_region(
    ws,
    df: pd.DataFrame,
    start_row: int,
    col_specs,
    table_name: Optional[str] = None,
    zebra: bool = True,
) -> int:
    """Write a styled DataFrame region; optionally wrap it in an Excel Table.

    Returns the next free row after the written block.
    """
    specs = [spec for spec in col_specs if spec[0] in df.columns]
    if not specs:
        return start_row

    header_row = start_row
    for col_idx, (_, header, _fmt) in enumerate(specs, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)
    _style_header_row(ws, header_row, len(specs))

    row_idx = header_row + 1
    for offset, (_, record) in enumerate(df.iterrows()):
        for col_idx, (col, _header, fmt) in enumerate(specs, start=1):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=_excel_value(record.get(col)),
            )
            cell.font = BODY_FONT
            cell.border = BORDER
            if fmt:
                cell.number_format = fmt
            if zebra and offset % 2 == 1:
                cell.fill = ALT_ROW_FILL
        row_idx += 1

    if table_name and row_idx - 1 > header_row:
        last_col = get_column_letter(len(specs))
        table = Table(
            displayName=table_name,
            ref=f"A{header_row}:{last_col}{row_idx - 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    return row_idx


class ExcelReportGenerator:
    """Generates a professional multi-sheet Excel report."""

    def __init__(self, output_config: Optional[OutputConfig] = None):
        """Initialize with output configuration (defaults to app config)."""
        self.output_config = output_config or config.output
        self.excel_dir = Path(self.output_config.excel_dir)
        self.report_name = self.output_config.report_name

    # ── Public API ──────────────────────────────────────────────

    def generate(
        self,
        df: pd.DataFrame,
        kpis: Dict[str, Any],
        metadata: Optional[Dict[str, Any]] = None,
    ) -> Path:
        """
        Build the Excel report workbook.

        Args:
            df: Cleaned product DataFrame.
            kpis: Output of ``KPIEngine.calculate_all()``.
            metadata: Optional info dict (api_source, record_count, ...).

        Returns:
            Path to the generated workbook.

        Raises:
            ExcelReportError: On empty data or missing KPI results.
        """
        logger.info("Generating Excel report")
        if df is None or len(df) == 0:
            raise ExcelReportError(
                "Cannot generate an Excel report from an empty dataset."
            )
        required_kpis = {"general", "category", "products"}
        if not isinstance(kpis, dict) or not required_kpis <= set(kpis):
            raise ExcelReportError(
                "kpis must contain 'general', 'category' and 'products' "
                "(use KPIEngine.calculate_all())."
            )

        now = datetime.now()
        meta = self._normalize_metadata(metadata or {}, df, now)

        try:
            self.excel_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise ExcelReportError(
                f"Cannot create output directory {self.excel_dir}: {exc}"
            ) from exc

        wb = Workbook()
        wb.properties.title = meta["report_title"]
        wb.properties.creator = "Automated Business Reporting Engine"

        dashboard_ws = wb.active
        dashboard_ws.title = "Dashboard"
        product_ws = wb.create_sheet("Product Data")
        category_ws = wb.create_sheet("Category Analysis")
        analysis_ws = wb.create_sheet("Product Analysis")
        metadata_ws = wb.create_sheet("Report Metadata")

        self._build_product_data(product_ws, df)
        category_last_row = self._build_category_analysis(category_ws, kpis)
        self._build_product_analysis(analysis_ws, kpis)
        self._build_dashboard(dashboard_ws, df, kpis, meta, now)
        self._add_category_charts(dashboard_ws, category_ws, category_last_row)
        self._build_metadata(metadata_ws, meta)

        wb.active = 0
        filepath = self.excel_dir / f"{self.report_name}_{now:%Y-%m-%d}.xlsx"
        try:
            wb.save(filepath)
        except Exception as exc:  # noqa: BLE001 - surface any save failure
            raise ExcelReportError(f"Failed to save Excel report: {exc}") from exc

        logger.info(f"Excel report saved to {filepath}")
        return filepath

    # ── Metadata ────────────────────────────────────────────────

    def _normalize_metadata(
        self, metadata: Dict[str, Any], df: pd.DataFrame, now: datetime
    ) -> Dict[str, Any]:
        meta = dict(metadata)
        meta.setdefault("report_title", self.report_name)
        meta.setdefault("report_version", REPORT_VERSION)
        meta.setdefault("api_source", str(config.api.url))
        meta.setdefault("record_count", int(len(df)))
        meta.setdefault("generation_timestamp", f"{now:%Y-%m-%d %H:%M:%S}")
        return meta

    # ── Sheet builders ──────────────────────────────────────────

    def _build_product_data(self, ws, df: pd.DataFrame) -> None:
        """Sheet 2 - full processed dataset as an Excel table."""
        _write_data_region(ws, df, 1, PRODUCT_COLUMNS, table_name="ProductData")
        _auto_widths(ws, df, PRODUCT_COLUMNS)
        ws.freeze_panes = "A2"
        logger.debug("Writing Product Data sheet")

    def _build_category_analysis(self, ws, kpis: Dict[str, Any]) -> int:
        """Sheet 3 - category metrics; returns last data row (for charts)."""
        category_df = kpis["category"]
        next_row = _write_data_region(
            ws, category_df, 1, CATEGORY_COLUMNS, table_name="CategoryAnalysis"
        )
        _auto_widths(ws, category_df, CATEGORY_COLUMNS)
        ws.freeze_panes = "A2"
        logger.debug("Writing Category Analysis sheet")
        return next_row - 1

    def _build_product_analysis(self, ws, kpis: Dict[str, Any]) -> None:
        """Sheet 4 - product-level ranking tables."""
        row = 1
        for key, title, specs in PRODUCT_ANALYSIS_SECTIONS:
            section_df = kpis["products"].get(key)
            if section_df is None or len(section_df) == 0:
                continue
            ws.cell(row=row, column=1, value=title).font = SECTION_FONT
            row += 1
            row = _write_data_region(
                ws, section_df, row, specs, table_name=TABLE_NAMES[key]
            )
            row += 1  # blank separator between sections
        ws.freeze_panes = "A4"
        logger.debug("Writing Product Analysis sheet")

    def _build_dashboard(
        self,
        ws,
        df: pd.DataFrame,
        kpis: Dict[str, Any],
        meta: Dict[str, Any],
        now: datetime,
    ) -> None:
        """Sheet 1 - title block, KPI summary, status and top products."""
        ncols = 9
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        title_cell = ws.cell(row=1, column=1, value=meta["report_title"])
        title_cell.font = TITLE_FONT
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        subtitle = ws.cell(
            row=2,
            column=1,
            value="Automated Business Reporting Engine - Catalog Overview",
        )
        subtitle.font = SUBTITLE_FONT
        subtitle.alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        generated = ws.cell(
            row=3,
            column=1,
            value=(
                f"Generated: {now:%Y-%m-%d %H:%M:%S}    |    "
                f"Data Source: {meta['api_source']}"
            ),
        )
        generated.font = BODY_FONT
        generated.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A4"

        row = self._build_kpi_table(ws, 5, kpis["general"])
        row = self._build_inventory_status(ws, row + 1, df)
        self._build_top_products(ws, row + 1, kpis)

        for col, width in {"A": 36, "B": 26, "C": 22, "D": 14, "E": 20}.items():
            ws.column_dimensions[col].width = width
        logger.debug("Writing Dashboard sheet")

    def _build_kpi_table(self, ws, start_row: int, general: Dict[str, Any]) -> int:
        """KPI summary label/value table on the Dashboard."""
        row = _write_section_title(ws, start_row, "KPI Summary")
        ws.cell(row=row, column=1, value="Metric")
        ws.cell(row=row, column=2, value="Value")
        _style_header_row(ws, row, 2)
        row += 1

        for offset, (key, label, fmt) in enumerate(KPI_META):
            ws.cell(row=row, column=1, value=label).font = BODY_FONT
            value_cell = ws.cell(row=row, column=2, value=_excel_value(general.get(key)))
            value_cell.font = BODY_FONT
            value_cell.number_format = fmt
            for col in (1, 2):
                cell = ws.cell(row=row, column=col)
                cell.border = BORDER
                if offset % 2 == 1:
                    cell.fill = ALT_ROW_FILL
            row += 1
        return row

    def _build_inventory_status(self, ws, start_row: int, df: pd.DataFrame) -> int:
        """Product count per stock status on the Dashboard."""
        if "stock_status" not in df.columns:
            return start_row
        row = _write_section_title(ws, start_row, "Inventory Status")
        ws.cell(row=row, column=1, value="Status")
        ws.cell(row=row, column=2, value="Product Count")
        _style_header_row(ws, row, 2)
        row += 1

        counts = df["stock_status"].value_counts().to_dict()
        for offset, status in enumerate(STOCK_STATUS_ORDER):
            if status not in counts:
                continue
            ws.cell(row=row, column=1, value=status).font = BODY_FONT
            count_cell = ws.cell(row=row, column=2, value=int(counts[status]))
            count_cell.font = BODY_FONT
            count_cell.number_format = INTEGER
            for col in (1, 2):
                cell = ws.cell(row=row, column=col)
                cell.border = BORDER
                if offset % 2 == 1:
                    cell.fill = ALT_ROW_FILL
            row += 1
        return row

    def _build_top_products(self, ws, start_row: int, kpis: Dict[str, Any]) -> int:
        """Top 5 products by inventory value on the Dashboard."""
        row = _write_section_title(
            ws, start_row, "Top Products by Inventory Value", max_col=5
        )
        top_df = kpis["products"].get("top_by_inventory_value")
        if top_df is None or len(top_df) == 0:
            return row
        specs = [
            ("title", "Product", None),
            ("category", "Category", None),
            ("price", "Price", CURRENCY),
            ("stock", "Stock", INTEGER),
            ("inventory_value", "Inventory Value", CURRENCY),
        ]
        return _write_data_region(ws, top_df.head(5), row, specs)

    def _add_category_charts(self, ws, category_ws, last_data_row: int) -> None:
        """Add category pie and inventory bar charts anchored on the Dashboard."""
        if last_data_row < 2:
            logger.warning("Not enough category data to build charts")
            return

        categories = Reference(
            category_ws, min_col=1, min_row=2, max_row=last_data_row
        )

        pie = PieChart()
        pie.title = "Products by Category"
        pie.add_data(
            Reference(
                category_ws,
                min_col=2,
                min_row=1,
                max_col=2,
                max_row=last_data_row,
            ),
            titles_from_data=True,
        )
        pie.set_categories(categories)
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = False
        pie.width = 16
        pie.height = 10
        ws.add_chart(pie, "F7")

        bar = BarChart()
        bar.type = "col"
        bar.title = "Inventory Value by Category"
        bar.add_data(
            Reference(
                category_ws,
                min_col=6,
                min_row=1,
                max_col=6,
                max_row=last_data_row,
            ),
            titles_from_data=True,
        )
        bar.set_categories(categories)
        bar.legend = None
        bar.y_axis.title = "Inventory Value"
        bar.width = 16
        bar.height = 10
        ws.add_chart(bar, "F25")
        logger.debug("Added category charts to Dashboard")

    def _build_metadata(self, ws, meta: Dict[str, Any]) -> None:
        """Sheet 5 - report/API metadata key-value table."""
        rows = [
            ("Report Name", meta["report_title"]),
            ("Report Version", meta["report_version"]),
            ("API Source", meta["api_source"]),
            ("Generation Timestamp", meta["generation_timestamp"]),
            ("Number of Records", meta["record_count"]),
            ("Report Type", "Product Catalog / Inventory Overview"),
        ]
        for row_idx, (label, value) in enumerate(rows, start=1):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = Font(bold=True)
            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = BODY_FONT
            for col in (1, 2):
                ws.cell(row=row_idx, column=col).border = BORDER
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 70
        logger.debug("Writing Report Metadata sheet")


def generate_excel_report(
    df: pd.DataFrame,
    kpis: Dict[str, Any],
    metadata: Optional[Dict[str, Any]] = None,
    output_config: Optional[OutputConfig] = None,
) -> Path:
    """Convenience wrapper for generating the Excel report."""
    return ExcelReportGenerator(output_config=output_config).generate(df, kpis, metadata)